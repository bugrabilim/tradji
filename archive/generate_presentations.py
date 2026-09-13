from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent
ART = ROOT / 'artifacts'
ART.mkdir(exist_ok=True)

NAVY = RGBColor(15, 23, 42)
BLUE = RGBColor(37, 99, 235)
GREEN = RGBColor(22, 163, 74)
ORANGE = RGBColor(245, 158, 11)
MUTED = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)
BORDER = RGBColor(226, 232, 240)


def set_wide(prs):
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)


def add_title(slide, text, subtitle=''):
    box = slide.shapes.add_textbox(Inches(.7), Inches(.4), Inches(12), Inches(.7))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        box = slide.shapes.add_textbox(Inches(.72), Inches(1.02), Inches(11.5), Inches(.4))
        p = box.text_frame.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(14)
        p.font.color.rgb = MUTED


def add_card(slide, x, y, w, h, heading, body, accent=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = BORDER
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(.08), Inches(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(x+.22), Inches(y+.18), Inches(w-.4), Inches(.45))
    p = box.text_frame.paragraphs[0]
    p.text = heading
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY
    box = slide.shapes.add_textbox(Inches(x+.22), Inches(y+.75), Inches(w-.4), Inches(h-.9))
    p = box.text_frame.paragraphs[0]
    p.text = body
    p.font.size = Pt(13)
    p.font.color.rgb = MUTED


def main_deck():
    prs = Presentation()
    set_wide(prs)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box = slide.shapes.add_textbox(Inches(.9), Inches(1.35), Inches(7), Inches(1))
    p = box.text_frame.paragraphs[0]
    p.text = 'Tradji'
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    box = slide.shapes.add_textbox(Inches(.95), Inches(2.6), Inches(7.4), Inches(1.3))
    p = box.text_frame.paragraphs[0]
    p.text = 'Kaza tespit surecini fotograf, konum ve akilli veri girisiyle sadelestiren dijital asistan'
    p.font.size = Pt(23)
    p.font.color.rgb = RGBColor(203, 213, 225)
    mockup = ART / 'Tradji_Mockup.png'
    if mockup.exists():
        slide.shapes.add_picture(str(mockup), Inches(8.7), Inches(.8), width=Inches(4.1))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Problem', 'Kaza ani zaten stresli. Tutanak sureci bu stresi buyutuyor.')
    add_card(slide, .8, 1.65, 3.7, 4.5, 'Manuel veri', 'Ruhsat, ehliyet, police ve arac bilgilerinin elle girilmesi.', ORANGE)
    add_card(slide, 4.8, 1.65, 3.7, 4.5, 'Adres ve kroki', 'Kullanici adresi bilmeyebilir; kroki cizmek zor ve hataya aciktir.', BLUE)
    add_card(slide, 8.8, 1.65, 3.7, 4.5, 'Eksik belge riski', 'Eksik veya tutarsiz bilgi sigorta surecini yavaslatabilir.', GREEN)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Cozum: Tradji', 'Yazmak yerine cek, konumu al, konus ve onayla.')
    if mockup.exists():
        slide.shapes.add_picture(str(mockup), Inches(.7), Inches(1.55), width=Inches(6.2))
    add_card(slide, 7.2, 1.7, 5.2, 1.1, 'Fotograf', 'Kaza ve belge fotograflari tek akista toplanir.')
    add_card(slide, 7.2, 3.1, 5.2, 1.1, 'Konum', 'GPS ile koordinat ve adres kolaylastirilir.', GREEN)
    add_card(slide, 7.2, 4.5, 5.2, 1.1, 'Tutanak', 'Veriler A4 veya PDF ciktisina donusturulur.', ORANGE)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Teknolojik Yaklasim')
    add_card(slide, .8, 1.55, 3.8, 2.0, 'OCR', 'Ruhsat, ehliyet, police ve plaka alanlarini fotograflardan okuma.')
    add_card(slide, 4.75, 1.55, 3.8, 2.0, 'Konum Servisleri', 'GPS ve reverse geocoding ile kaza noktasini otomatiklestirme.', GREEN)
    add_card(slide, 8.7, 1.55, 3.8, 2.0, 'Ses -> Metin', 'Surucu beyanini konusarak olusturma.', ORANGE)
    add_card(slide, 2.8, 4.0, 3.8, 2.0, 'Goruntu Isleme', 'Arac pozisyonu ve gelecekte otomatik kroki.')
    add_card(slide, 6.75, 4.0, 3.8, 2.0, 'PDF / A4', 'Standart tutanagi yazdirilabilir belgeye donusturme.', GREEN)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'MVP Stratejisi', 'Once entegrasyonsuz deger, sonra kurum entegrasyonlari')
    add_card(slide, .9, 1.7, 5.7, 4.5, 'FAZ 1 - Kapali Devre MVP', 'Form ve kullanici akisi\nFotograf / belge yukleme\nKonum\nBasit kroki\nSesli beyan\nPDF / A4 cikti', BLUE)
    add_card(slide, 6.85, 1.7, 5.55, 4.5, 'FAZ 2 - Entegrasyonlar', 'Sigorta bildirimi\nHasar dosyasi acma\nYetkili kamu verisi\nGercek zamanli olay paylasimi\nDosya durum takibi', ORANGE)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, 'Deger Onerisi')
    box = slide.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11), Inches(1.5))
    p = box.text_frame.paragraphs[0]
    p.text = 'Daha az yazi. Daha az hata. Daha hizli tutanak.'
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    add_card(slide, 1.2, 3.4, 3.3, 2.1, 'Surucu', 'Stresli anda basit ve yonlendirilmis deneyim.', BLUE)
    add_card(slide, 5.0, 3.4, 3.3, 2.1, 'Sigorta', 'Daha standart ve eksiksiz veri potansiyeli.', GREEN)
    add_card(slide, 8.8, 3.4, 3.3, 2.1, 'Ekosistem', 'Kaza anindan hasar dosyasina uzanan dijital akis.', ORANGE)

    prs.save(ART / 'Tradji_Kaza_Tespit_Sunumu.pptx')


def tech_deck():
    prs = Presentation()
    set_wide(prs)
    slides = [
        ('Kullanici Merkezli Tasarim + Agile', [
            ('Kullanici Merkezli Tasarim', 'Kolay, anlasilir ve dusuk bilissel yuk.', BLUE),
            ('Agile', 'Kucuk iterasyonlarda calisan prototipler.', GREEN),
        ]),
        ('RESTful API + OCR', [
            ('RESTful API', 'Frontend ve backend ayrimi; entegrasyona hazir mimari.', BLUE),
            ('OCR', 'Ruhsat, ehliyet, police ve plaka bilgisini fotograflardan cikarma.', ORANGE),
        ]),
        ('Konum + PDF + Ses Tanima', [
            ('Konum Servisleri', 'GPS ve harita servisleri.', GREEN),
            ('PDF / Gorsellestirme', 'Tutanagi A4 ciktisina donusturme.', BLUE),
            ('Ses Tanima', 'Beyani konusmadan metne cevirme.', ORANGE),
        ]),
    ]
    for heading, cards in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_title(slide, heading, 'Tradji - Yontemler ve teknolojik yaklasimlar')
        gap = .35
        total = 11.8
        width = (total - gap * (len(cards)-1)) / len(cards)
        x = .75
        for h, b, a in cards:
            add_card(slide, x, 1.75, width, 4.8, h, b, a)
            x += width + gap
    prs.save(ART / 'Tradji_Teknolojik_Yaklasimlar_Sunumu.pptx')


if __name__ == '__main__':
    main_deck()
    tech_deck()
